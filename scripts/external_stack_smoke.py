"""Exercise the deployed external stack through the public HTTP API.

The script intentionally avoids FastAPI TestClient and SQLite. It proves a real
PostgreSQL/Redis/Milvus/OpenSearch/Celery deployment by uploading a document,
running hybrid retrieval, executing a Celery evaluation run, and cleaning up.
"""

from __future__ import annotations

import json
import os
import sys
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import httpx


API_BASE_URL = os.getenv("OA_RAG_API_URL", "http://127.0.0.1:18080").rstrip("/")
INTERNAL_TOKEN = os.getenv("INTERNAL_TOKEN", "integration-internal-token")
POLL_TIMEOUT_SECONDS = int(os.getenv("OA_RAG_SMOKE_TIMEOUT_SECONDS", "180"))
POLL_INTERVAL_SECONDS = float(os.getenv("OA_RAG_SMOKE_POLL_INTERVAL_SECONDS", "2"))


def build_docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("调岗审批流程联调样本", level=1)
    document.add_paragraph("制度编号：HR-TRF-INTEGRATION-01。")
    document.add_paragraph("调岗申请必须经过两级审批：部门负责人审批、HR 系统负责人审批。")
    document.add_paragraph("两级审批完成后，申请状态更新为已完成。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def request_json(client: httpx.Client, method: str, path: str, **kwargs) -> dict:
    response = client.request(method, f"{API_BASE_URL}{path}", **kwargs)
    try:
        body = response.json()
    except ValueError:
        body = {"raw": response.text[:500]}
    if response.status_code >= 400:
        raise RuntimeError(f"{method} {path} failed: HTTP {response.status_code}; {body}")
    return body


def wait_for_job(client: httpx.Client, job_uuid: str) -> dict:
    deadline = time.monotonic() + POLL_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        body = request_json(client, "GET", f"/api/v1/jobs/{job_uuid}")
        status = body.get("data", {}).get("status")
        if status == "completed":
            return body
        if status == "failed":
            raise RuntimeError(f"ingestion job failed: {body}")
        time.sleep(POLL_INTERVAL_SECONDS)
    raise TimeoutError(f"ingestion job {job_uuid} did not complete within {POLL_TIMEOUT_SECONDS}s")


def wait_for_run(client: httpx.Client, run_uuid: str) -> dict:
    deadline = time.monotonic() + POLL_TIMEOUT_SECONDS
    while time.monotonic() < deadline:
        body = request_json(client, "GET", f"/api/v1/evaluation/runs/{run_uuid}")
        status = body.get("data", {}).get("status")
        if status == "completed":
            return body
        if status == "failed":
            raise RuntimeError(f"evaluation run failed: {body}")
        time.sleep(POLL_INTERVAL_SECONDS)
    raise TimeoutError(f"evaluation run {run_uuid} did not complete within {POLL_TIMEOUT_SECONDS}s")


def write_artifact(result: dict) -> Path:
    root = Path(__file__).resolve().parents[1]
    output_dir = root / "qa-outputs" / "external-stack" / datetime.now().strftime("%Y%m%d-%H%M%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "smoke-result.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main() -> int:
    smoke_id = uuid4().hex[:10]
    document_title = f"调岗审批流程联调样本-{smoke_id}"
    result: dict = {"smoke_id": smoke_id, "api_base_url": API_BASE_URL, "checks": {}}
    doc_uuid: str | None = None
    run_uuid: str | None = None
    exit_code = 0

    with httpx.Client(timeout=60.0) as client:
        try:
            health = request_json(client, "GET", "/api/v1/health")
            health_data = health.get("data", {})
            result["checks"]["health"] = health_data
            if health_data.get("postgres") != "up" or health_data.get("redis") != "up":
                raise RuntimeError(f"PostgreSQL or Redis is unavailable: {health_data}")
            if health_data.get("probes", {}).get("zilliz") != "up":
                raise RuntimeError(f"Milvus probe is unavailable: {health_data}")

            upload = request_json(
                client,
                "POST",
                "/api/v1/documents/upload",
                data={
                    "title": document_title,
                    "source_type": "rule_doc",
                    "source_module": "hr",
                    "owner_dept": "hr",
                    "version": smoke_id,
                },
                files={
                    "file": (
                        f"transfer-policy-{smoke_id}.docx",
                        build_docx_bytes(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            doc_uuid = upload["data"]["doc_uuid"]
            job_uuid = upload["data"]["job_uuid"]
            result["checks"]["upload"] = {"doc_uuid": doc_uuid, "job_uuid": job_uuid}
            result["checks"]["ingestion"] = wait_for_job(client, job_uuid).get("data", {})

            debug_search = request_json(
                client,
                "POST",
                "/api/v1/retrieval/debug-search",
                json={
                    "query": "HR-TRF-INTEGRATION-01 调岗 两级审批",
                    "top_k": 3,
                    "strategy": "hybrid",
                    "filters": {"source_module": ["hr"]},
                    "user_context": {"user_id": "external-stack-smoke"},
                },
            )
            debug_data = debug_search.get("data", {})
            if not debug_data.get("dense_hits") or not debug_data.get("sparse_hits"):
                raise RuntimeError(f"Hybrid did not receive both providers: {debug_data}")
            result["checks"]["hybrid_search"] = {
                "dense_hit_count": len(debug_data["dense_hits"]),
                "sparse_hit_count": len(debug_data["sparse_hits"]),
                "top_title": debug_data.get("hits", [{}])[0].get("title"),
            }

            agent_answer = request_json(
                client,
                "POST",
                "/api/v1/qa/agent-answer",
                headers={"X-Internal-Token": INTERNAL_TOKEN},
                json={
                    "question": "调岗申请需要哪两级审批？",
                    "top_k": 3,
                    "retrieval_strategy": "hybrid",
                    "filters": {"source_module": ["hr"]},
                    "user_context": {"user_id": "external-stack-smoke"},
                },
            )
            agent_data = agent_answer.get("data", {})
            if agent_data.get("answer_status") != "grounded" or not agent_data.get("citations"):
                raise RuntimeError(f"Agent did not return grounded evidence: {agent_data}")
            result["checks"]["agent"] = {
                "answer_status": agent_data["answer_status"],
                "citation_count": len(agent_data["citations"]),
                "trace": agent_data.get("trace", []),
            }

            dataset = request_json(
                client,
                "POST",
                "/api/v1/evaluation/datasets",
                json={
                    "name": f"external-stack-smoke-{smoke_id}",
                    "description": "External PostgreSQL/Redis/Milvus/OpenSearch smoke test",
                    "queries": [
                        {
                            "query_text": "调岗申请两级审批",
                            "expected_doc_titles": [document_title],
                            "expected_terms": ["部门负责人", "HR 系统负责人"],
                        }
                    ],
                },
            )
            dataset_uuid = dataset["data"]["dataset_uuid"]
            run = request_json(
                client,
                "POST",
                "/api/v1/evaluation/runs",
                json={"dataset_uuid": dataset_uuid, "doc_uuids": [doc_uuid], "retrieval_strategy": "dense"},
            )
            run_uuid = run["data"]["run_uuid"]
            if run["data"]["status"] != "pending":
                raise RuntimeError(f"Celery evaluation should initially be pending: {run}")
            run_detail = wait_for_run(client, run_uuid).get("data", {})
            result["checks"]["celery_evaluation"] = {
                "run_uuid": run_uuid,
                "status": run_detail.get("status"),
                "summary": run_detail.get("summary"),
            }
        except Exception as exc:
            result["error"] = str(exc)
            exit_code = 1
        finally:
            if run_uuid:
                try:
                    request_json(client, "DELETE", f"/api/v1/evaluation/runs/{run_uuid}")
                    result["checks"]["evaluation_cleanup"] = "completed"
                except Exception as exc:
                    result["checks"]["evaluation_cleanup"] = f"failed: {exc}"
            if doc_uuid:
                try:
                    request_json(client, "DELETE", f"/api/v1/documents/{doc_uuid}")
                    result["checks"]["document_cleanup"] = "completed"
                except Exception as exc:
                    result["checks"]["document_cleanup"] = f"failed: {exc}"

    result["artifact_path"] = str(write_artifact(result))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
