"""检索与 QA 端到端测试：验证搜索、问答、权限和评测的完整流程。

测试覆盖：
- 授权用户检索返回结果
- QA 回答返回带引用的 grounded 回答
- 内部 Token 验证拒绝非法请求
- 调试搜索返回评分分解
- 高阈值导致 insufficient_evidence
- 中文自然语言查询支持
- 过滤条件大小写不敏感
- 评测运行的数据隔离
"""

# 启用 PEP 563 延迟注解求值，避免循环引用并提升模块导入速度
from __future__ import annotations

import fitz  # PyMuPDF：用于动态生成测试用 PDF 文档

def _upload_sample_document(client, sample_pdf_bytes):
    """上传示例 PDF 文档并返回 doc_uuid。

    Args:
        client: FastAPI TestClient 实例
        sample_pdf_bytes: conftest.py 中定义的 fixture，预生成的 PDF 二进制数据
    """
    response = client.post(
        "/api/v1/documents/upload",
        data={"source_type": "rule_doc", "source_module": "hr", "owner_dept": "hr"},
        files={"file": ("policy.pdf", sample_pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()["data"]["doc_uuid"]


def _upload_custom_pdf_document(
    client,
    *,
    title: str,
    text: str,
    source_module: str = "hr",
) -> str:
    """使用 fitz (PyMuPDF) 动态创建包含指定文本的 PDF 并上传。

    面试点：为什么测试中需要动态生成 PDF？
    答：这样可以为每个测试用例创建特定内容的文档（如包含特定标题、关键词），
    使测试具有可复现性和独立性，不依赖外部文件。

    Args:
        client: TestClient 实例
        title: 文档标题
        text: PDF 中的文本内容
        source_module: 来源模块（默认 hr）
    """
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)  # 在 (72, 72) 坐标处插入文本
    payload = document.tobytes()  # 转为 PDF 字节流
    document.close()
    response = client.post(
        "/api/v1/documents/upload",
        data={
            "source_type": "rule_doc",
            "source_module": source_module,
            "owner_dept": "hr",
            "title": title,
        },
        files={"file": (f"{title}.pdf", payload, "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()["data"]["doc_uuid"]

def _upload_document_with_resource_meta(
    client,
    sample_pdf_bytes,
    *,
    extra_meta: str,
    tags: str | None = None,
    version: str = "v1",
    file_name: str = "resource_policy.pdf",
):
    """上传带有扩展元数据（extra_meta、tags、version）的文档。

    用于测试文档的扩展元数据字段是否正确存储和检索。
    """
    data = {
        "source_type": "rule_doc",
        "source_module": "hr",
        "owner_dept": "hr",
        "version": version,
        "extra_meta": extra_meta,
    }
    if tags is not None:
        data["tags"] = tags
    response = client.post(
        "/api/v1/documents/upload",
        data=data,
        files={"file": (file_name, sample_pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()["data"]["doc_uuid"]

def _upload_document_without_owner(client, sample_pdf_bytes, *, access_level: str):
    """上传不带 owner_dept 的文档（用于测试权限边界情况）。"""
    response = client.post(
        "/api/v1/documents/upload",
        data={"source_type": "rule_doc", "source_module": "hr", "access_level": access_level},
        files={"file": (f"{access_level}_policy.pdf", sample_pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()["data"]["doc_uuid"]

def _upload_sample_docx_document(client, sample_docx_bytes):
    """上传示例 DOCX 文档（用于测试中文查询和多种文件格式）。"""
    response = client.post(
        "/api/v1/documents/upload",
        data={"source_type": "rule_doc", "source_module": "hr", "owner_dept": "hr"},
        files={
            "file": (
                "policy.docx",
                sample_docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    return response.json()["data"]["doc_uuid"]

def _create_identity_user(
    client,
    *,
    user_id: str = "u1001",
    departments: list[str] | None = None,
    roles: list[str] | None = None,
    external_source: str | None = None,
    external_id: str | None = None,
):
    """创建身份用户：先创建部门和角色，再创建用户并关联。

    面试点：为什么 assert 接受 200 或 409？
    答：409 (Conflict) 表示资源已存在（如部门已创建），这在重复运行测试时常见。
    接受 409 保证测试的幂等性——多次运行不因资源冲突而失败。

    Args:
        client: TestClient 实例
        user_id: 用户 ID（默认 u1001）
        departments: 部门列表（默认 ["hr"]）
        roles: 角色列表（默认 ["employee"]）
        external_source: 外部来源（如 LDAP）
        external_id: 外部 ID
    """
    departments = departments or ["hr"]
    roles = roles or ["employee"]
    # 创建部门（幂等：已存在则返回 409）
    for dept_code in departments:
        response = client.post(
            "/api/v1/admin/departments",
            json={"dept_code": dept_code, "dept_name": dept_code},
        )
        assert response.status_code in {200, 409}
    # 创建角色（幂等）
    for role_code in roles:
        response = client.post(
            "/api/v1/admin/roles",
            json={"role_code": role_code, "role_name": role_code},
        )
        assert response.status_code in {200, 409}
    # 创建用户
    response = client.post(
        "/api/v1/admin/users",
        json={
            "user_id": user_id,
            "display_name": user_id,
            "department_codes": departments,
            "role_codes": roles,
            "primary_dept_code": departments[0] if departments else None,
            "external_source": external_source,
            "external_id": external_id,
        },
    )
    assert response.status_code in {200, 409}

def _user_context(user_id: str = "u1001") -> dict:
    """构建用户上下文字典（用于检索请求）。"""
    return {"user_id": user_id}

def test_search_returns_hits_for_authorized_user(client, sample_pdf_bytes):
    """测试授权用户检索返回命中结果。

    面试点：为什么测试中要切换 INGESTION_MODE？
    答：默认配置可能为 async（异步摄取），但测试需要同步等待摄取完成
    才能确保文档已被索引。切换到 sync 模式保证测试的确定性。
    """
    # 确保使用同步摄取模式
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()  # 清除配置缓存，使环境变量生效

    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/retrieval/search",
        json={
            "query": "transfer approval two levels",
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == 0  # 成功状态码
    assert len(payload["data"]["hits"]) >= 1  # 至少一个命中
    assert "vector_score" in payload["data"]["hits"][0]  # 包含向量得分
    assert "text_score" in payload["data"]["hits"][0]  # 包含文本得分

    # 恢复异步模式
    os.environ["INGESTION_MODE"] = "async"
    reset_settings_cache()


def test_search_without_filters_uses_index_candidates(client, sample_pdf_bytes):
    """无筛选且无权限插件时，检索应正常命中而不依赖全量文档预加载。"""
    import os

    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache

    reset_settings_cache()
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/retrieval/search",
        json={"query": "transfer approval two levels", "top_k": 5},
    )

    assert response.status_code == 200
    assert response.json()["data"]["hits"]

    os.environ["INGESTION_MODE"] = "async"
    reset_settings_cache()

def test_qa_answer_returns_grounded_citations(client, sample_pdf_bytes):
    """测试 QA 回答返回 grounded 状态和引用片段。

    验证点：
    - 回答状态为 grounded（证据充分）
    - 回答中包含"基于命中的引用片段"（使用了 fallback 模板）
    - citations 列表非空
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()

    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/qa/answer",
        headers={"X-Internal-Token": "test-internal-token"},  # 内部 Token
        json={
            "question": "transfer approval two levels",
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
            "generation_options": {"temperature": 0.1, "max_tokens": 300},
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == 0
    assert payload["data"]["citations"]  # 有引用
    assert payload["data"]["answer_status"] == "grounded"  # 证据充分
    assert "基于命中的引用片段" in payload["data"]["answer"]  # 使用了模板回答

    os.environ["INGESTION_MODE"] = "async"
    reset_settings_cache()

def test_qa_answer_rejects_bad_internal_token(client, sample_pdf_bytes):
    """测试错误的内部 Token 被拒绝（返回 401）。

    面试点：为什么内部 Token 验证返回 401 而非 403？
    答：401 表示认证失败（未提供有效凭证），403 表示授权失败（已认证但无权限）。
    错误的 Token 属于认证问题，应返回 401。
    """
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/qa/answer",
        headers={"X-Internal-Token": "wrong-token"},  # 故意使用错误 Token
        json={
            "question": "transfer approval two levels",
            "top_k": 5,
            "user_context": {
                "user_id": "u1001",
                "roles": ["hr_admin"],
                "departments": ["hr"],
                "is_super_admin": False,
            },
        },
    )
    assert response.status_code == 401  # 认证失败
    payload = response.json()
    assert payload["code"] == 40101  # AUTHENTICATION_FAILED 错误码

def test_debug_search_returns_score_breakdown(client, sample_pdf_bytes):
    """测试调试搜索返回评分分解详情。

    验证点：
    - ranking_debug 列表非空
    - 每条调试信息包含 vector_score 和 text_score
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/retrieval/debug-search",
        json={
            "query": "transfer approval two levels",
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["data"]["ranking_debug"]  # 有评分分解
    assert "vector_score" in payload["data"]["ranking_debug"][0]
    assert "text_score" in payload["data"]["ranking_debug"][0]

def test_qa_answer_returns_insufficient_evidence_for_high_threshold(client, sample_pdf_bytes):
    """测试当查询与文档无关时，QA 返回 insufficient_evidence。

    使用 "completely unrelated cafeteria menu" 作为查询，
    确保检索得分低于证据阈值，触发 insufficient_evidence 状态。
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_sample_document(client, sample_pdf_bytes)
    response = client.post(
        "/api/v1/qa/answer",
        headers={"X-Internal-Token": "test-internal-token"},
        json={
            "question": "completely unrelated cafeteria menu",  # 与 HR 文档无关的查询
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["data"]["answer_status"] == "insufficient_evidence"  # 证据不足
    assert "未找到足够依据" in payload["data"]["answer"]  # 拒绝回答模板

def test_search_supports_natural_chinese_query(client, sample_docx_bytes):
    """测试中文自然语言查询支持。

    面试点：text_score > 0 说明什么？
    答：说明 BM25 稀疏检索对 jieba 中文分词后的关键词匹配生效了，
    证明词级 tokenization 在正常工作。
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_sample_docx_document(client, sample_docx_bytes)
    response = client.post(
        "/api/v1/retrieval/search",
        json={
            "query": "调岗申请的二级审批是怎么要求的？",  # 中文查询
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == 0
    assert payload["data"]["hits"]  # 有命中
    assert payload["data"]["hits"][0]["text_score"] > 0  # BM25 关键词得分 > 0

def test_search_source_module_filter_is_case_insensitive(client, sample_docx_bytes):
    """测试 source_module 过滤大小写不敏感。

    文档以 source_module="OA"（大写）上传，搜索时用 "oa"（小写）过滤，
    应能命中。这验证了 _load_documents 中的 func.lower() 处理。
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    # 上传大写 source_module 的文档
    response = client.post(
        "/api/v1/documents/upload",
        data={"source_type": "rule_doc", "source_module": "OA", "owner_dept": "hr"},
        files={
            "file": (
                "policy.docx",
                sample_docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200

    # 用小写 "oa" 搜索，应能命中
    response = client.post(
        "/api/v1/retrieval/search",
        json={
            "query": "调岗申请的二级审批是怎么要求的？",
            "top_k": 5,
            "filters": {"source_module": ["oa"]},  # 小写过滤
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == 0
    assert payload["data"]["hits"]  # 有命中 → 大小写不敏感生效

def test_qa_answer_supports_natural_chinese_question(client):
    """测试 QA 支持中文自然语言问题。

    验证完整的 QA 流程：中文问题 → 检索 → 生成回答 → grounded 状态。
    """
    import os
    from io import BytesIO
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache
    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    # 使用不含表格的单块 DOCX：测试环境用 SHA256 确定性哈希向量（无语义能力），
    # 若文档含表格，表格块会被误排到答案块之上（向量分高但无关键词），证据门控判不足。
    # 真实语义 Embedding 环境下含表格的原文同样能 grounded，这里单块化仅保证测试确定性。
    from docx import Document as DocxDocument
    document = DocxDocument()
    document.add_heading("调岗规则", level=1)
    document.add_paragraph("调岗申请需要增加二级审批。")
    buf = BytesIO()
    document.save(buf)
    _upload_sample_docx_document(client, buf.getvalue())
    response = client.post(
        "/api/v1/qa/answer",
        headers={"X-Internal-Token": "test-internal-token"},
        json={
            "question": "调岗申请的二级审批是怎么要求的？",  # 中文问题
            "top_k": 5,
            "filters": {"source_module": ["hr"]},
            "user_context": _user_context("u1001"),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["code"] == 0
    assert payload["data"]["answer_status"] == "grounded"  # 证据充分
    assert payload["data"]["citations"]  # 有引用

def _create_evaluation_dataset(client) -> str:
    """辅助函数：创建评测数据集并返回 dataset_uuid。"""
    response = client.post(
        "/api/v1/evaluation/datasets",
        json={
            "name": "permission-aware-eval",
            "queries": [
                {
                    "query_text": "transfer approval two levels",
                    "expected_terms": ["approval"],
                }
            ],
        },
    )
    assert response.status_code == 200
    return response.json()["data"]["dataset_uuid"]

def _run_evaluation(client, dataset_uuid: str) -> dict:
    """辅助函数：运行评测并返回详情。

    面试点：为什么 assert run["status"] == "completed"？
    答：sync 模式下评测在请求线程中同步执行，创建运行后应立即完成。
    如果状态不是 completed，说明评测执行失败。
    """
    response = client.post(
        "/api/v1/evaluation/runs",
        json={
            "dataset_uuid": dataset_uuid,
            "retrieval_strategy": "dense",
        },
    )
    assert response.status_code == 200
    run = response.json()["data"]
    assert run["status"] == "completed"  # 同步模式下应立即完成

    detail_response = client.get(f"/api/v1/evaluation/runs/{run['run_uuid']}")
    assert detail_response.status_code == 200
    return detail_response.json()["data"]


def test_evaluation_run_uses_isolated_vector_collection_only(client):
    """测试评测运行的数据隔离：验证隔离文档不会被生产数据污染。

    测试场景：
    1. 上传 "Production Match" 文档（包含 transfer approval 相关内容）
    2. 上传 "Isolated Eval Document" 文档（仅包含 travel expense 内容）
    3. 创建评测运行，仅使用 "Isolated Eval Document" 作为 doc_uuids
    4. 验证：对 transfer approval 查询，Hit@1/3/5 均为 False，MRR=0
       → 证明评测确实隔离了数据，只搜索了指定的隔离文档

    面试点：这个测试证明了什么？
    答：证明了评测的数据隔离设计有效。如果评测使用了生产数据（包含
    "Production Match"），Hit@1 应该为 True。但测试结果显示全部 False，
    说明评测只在隔离的 evaluation_chunks 中搜索，不污染也不依赖生产数据。
    """
    import os
    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache

    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    # 上传生产文档（包含 transfer approval 内容）
    _upload_custom_pdf_document(
        client,
        title="Production Match",
        text="Transfer approval now requires two levels.",
    )
    # 上传隔离文档（仅包含 travel expense 内容）
    isolated_doc_uuid = _upload_custom_pdf_document(
        client,
        title="Isolated Eval Document",
        text="Travel expense claims require receipt attachment.",
    )

    # 创建评测数据集：期望命中 "Production Match"
    dataset_response = client.post(
        "/api/v1/evaluation/datasets",
        json={
            "name": "isolated-eval-only",
            "queries": [
                {
                    "query_text": "transfer approval two levels",
                    "expected_doc_titles": ["Production Match"],
                    "expected_terms": ["approval"],
                }
            ],
        },
    )
    assert dataset_response.status_code == 200
    dataset_uuid = dataset_response.json()["data"]["dataset_uuid"]

    # 创建评测运行：仅使用隔离文档
    run_response = client.post(
        "/api/v1/evaluation/runs",
        json={
            "dataset_uuid": dataset_uuid,
            "doc_uuids": [isolated_doc_uuid],  # 只包含隔离文档
            "retrieval_strategy": "dense",
        },
    )
    assert run_response.status_code == 200
    run = run_response.json()["data"]
    assert run["status"] == "completed"

    detail_response = client.get(f"/api/v1/evaluation/runs/{run['run_uuid']}")
    assert detail_response.status_code == 200
    result = detail_response.json()["data"]["results"][0]

    # 验证隔离性：隔离文档中不含 transfer approval 内容，所有指标应为 0
    assert result["hit_at_1"] is False   # 第一位未命中
    assert result["hit_at_3"] is False   # 前三位未命中
    assert result["hit_at_5"] is False   # 前五位未命中
    assert result["mrr"] == 0.0          # MRR 为 0
    assert result["expected_term_hit_rate"] == 0.0  # 关键词命中率为 0
    # 验证命中列表中不包含 "Production Match"（生产文档被隔离）
    assert all(hit["title"] != "Production Match" for hit in result["top_hits"])

    os.environ["INGESTION_MODE"] = "async"
    reset_settings_cache()


def test_evaluation_run_hybrid_uses_isolated_sparse_index(client):
    """Hybrid evaluation must use a Run-scoped BM25 index instead of dense fallback."""
    import os

    os.environ["INGESTION_MODE"] = "sync"
    from app.core.config import reset_settings_cache

    reset_settings_cache()
    _create_identity_user(client, user_id="u1001", departments=["hr"])
    _upload_custom_pdf_document(
        client,
        title="Hybrid Evaluation Policy",
        text="Transfer approval requires manager approval and HR approval.",
    )
    dataset_response = client.post(
        "/api/v1/evaluation/datasets",
        json={
            "name": "hybrid-isolated-evaluation",
            "queries": [
                {
                    "query_text": "transfer approval HR",
                    "expected_doc_titles": ["Hybrid Evaluation Policy"],
                    "expected_terms": ["approval"],
                }
            ],
        },
    )
    assert dataset_response.status_code == 200

    run_response = client.post(
        "/api/v1/evaluation/runs",
        json={
            "dataset_uuid": dataset_response.json()["data"]["dataset_uuid"],
            "retrieval_strategy": "hybrid",
            "retrieval_params": {"rrf_k": 60},
        },
    )
    assert run_response.status_code == 200
    run = run_response.json()["data"]
    assert run["status"] == "completed"

    detail_response = client.get(f"/api/v1/evaluation/runs/{run['run_uuid']}")
    assert detail_response.status_code == 200
    result = detail_response.json()["data"]["results"][0]
    assert result["debug_info"]["executed_strategy"] == "hybrid"
    assert result["top_hits"]
    assert result["top_hits"][0]["sparse_score"] is not None

    os.environ["INGESTION_MODE"] = "async"
    reset_settings_cache()
