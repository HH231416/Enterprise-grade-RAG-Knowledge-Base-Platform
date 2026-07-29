from __future__ import annotations

import base64
from pathlib import Path
from types import SimpleNamespace

import fitz
from unittest.mock import AsyncMock, MagicMock
from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image

from app.ingestion.parsers.docx_parser import DocxParser
from app.ingestion.parsers.html_parser import HtmlParser
from app.ingestion.parsers.pdf_parser import PdfWithImagesParser
from app.ingestion.parsers.xlsx_parser import XlsxParser
from app.ingestion.image_assets import ImageAsset
from app.ingestion.types import ParsedBlock
from app.integrations.image_caption import ImageCaptionProvider
from app.integrations.image_recognition import ImageRecognitionProvider, RecognitionResult
from app.ingestion.parsers.image_parser import ImageParser
from app.integrations.embedding import EmbeddingResult
from app.services.indexing import IndexingService


PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _fake_parse_assets(monkeypatch, captured):
    def parse_assets(self, assets):
        captured.extend(assets)
        return [
            ParsedBlock(
                text=f"image:{asset.source_name}",
                page_no=asset.page_no,
                sheet_name=asset.sheet_name,
                metadata=asset.metadata,
            )
            for asset in assets
        ]

    monkeypatch.setattr("app.ingestion.parsers.image_parser.ImageParser.parse_assets", parse_assets)


def test_pdf_embedded_images_keep_page_location(tmp_path, monkeypatch):
    pdf_path = tmp_path / "with-image.pdf"
    document = fitz.open()
    document.new_page()
    document.new_page().insert_image(fitz.Rect(10, 10, 30, 30), stream=PNG_BYTES)
    document.save(pdf_path)
    document.close()

    captured = []
    _fake_parse_assets(monkeypatch, captured)
    blocks = PdfWithImagesParser().parse(pdf_path)

    assert captured[0].page_no == 2
    assert any(block.page_no == 2 for block in blocks if block.text.startswith("image:"))


def test_docx_embedded_images_are_extracted(tmp_path, monkeypatch):
    image_path = tmp_path / "image.png"
    image_path.write_bytes(PNG_BYTES)
    docx_path = tmp_path / "with-image.docx"
    document = DocxDocument()
    document.add_paragraph("正文")
    document.add_picture(str(image_path))
    document.save(docx_path)

    captured = []
    _fake_parse_assets(monkeypatch, captured)
    blocks = DocxParser().parse(docx_path)

    assert len(captured) == 1
    assert captured[0].metadata["extraction_method"] == "docx_relationship"
    assert any(block.text == "正文" for block in blocks)


def test_xlsx_embedded_images_are_extracted(tmp_path, monkeypatch):
    image_path = tmp_path / "image.png"
    Image.new("RGB", (2, 2), color="white").save(image_path)
    xlsx_path = tmp_path / "with-image.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.add_image(XlsxImage(str(image_path)), "B3")
    workbook.save(xlsx_path)
    workbook.close()

    captured = []
    _fake_parse_assets(monkeypatch, captured)
    XlsxParser().parse(xlsx_path)

    assert len(captured) == 1
    assert captured[0].sheet_name == "Sheet1"
    assert captured[0].metadata["anchor_row"] == 3
    assert captured[0].metadata["anchor_column"] == 2


def test_xlsx_chart_data_is_searchable(tmp_path):
    xlsx_path = tmp_path / "with-chart.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sales"
    worksheet.append(["Quarter", "Revenue"])
    worksheet.append(["Q1", 10])
    worksheet.append(["Q2", 20])
    chart = BarChart()
    chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=3))
    worksheet.add_chart(chart, "D2")
    workbook.save(xlsx_path)
    workbook.close()

    blocks = XlsxParser().parse(xlsx_path)
    chart_block = next(block for block in blocks if block.chunk_type == "image_chart")

    assert chart_block.metadata["chart_series"][0]["categories"] == ["Q1", "Q2"]
    assert chart_block.metadata["chart_series"][0]["values"] == ["10", "20"]


def test_html_local_and_data_images_are_extracted(tmp_path):
    local_image = tmp_path / "local.png"
    local_image.write_bytes(PNG_BYTES)
    html_path = tmp_path / "page.html"
    data_uri = "data:image/png;base64," + base64.b64encode(PNG_BYTES).decode()
    html_path.write_text(
        f'<p>正文</p><img src="local.png" alt="本地图"><img src="{data_uri}" alt="内嵌图">',
        encoding="utf-8",
    )

    parser = HtmlParser()
    html_parser = parser  # keep the test focused on the safe extraction helper
    assets = html_parser._extract_local_images(
        [("local.png", "本地图"), (data_uri, "内嵌图"), ("https://example.com/a.png", "远程")],
        html_path,
    )

    assert len(assets) == 2
    assert assets[0].metadata["alt"] == "本地图"
    assert assets[1].metadata["extraction_method"] == "html_img"


def test_low_confidence_ocr_triggers_vlm_and_cache(monkeypatch):
    ocr = MagicMock(spec=ImageRecognitionProvider)
    ocr.is_configured.return_value = True
    ocr.recognize_image = AsyncMock(
        return_value=RecognitionResult(text="流程图", confidence=0.4, provider_name="ocr")
    )
    caption = MagicMock(spec=ImageCaptionProvider)
    caption.is_configured.return_value = True
    caption.caption = AsyncMock(return_value="流程图描述")
    parser = ImageParser(recognition_provider=ocr, caption_provider=caption)
    parser._cache_enabled = True
    parser._recognition_cache.clear()
    asset = ImageAsset(content=b"same-image", source_name="a.pdf#image-1", page_no=8)

    first = parser.parse_asset(asset)
    second = parser.parse_asset(asset)

    assert "流程图描述" in first[0].text
    assert first[0].page_no == 8
    assert second[0].metadata["cache_hit"] is True
    assert ocr.recognize_image.await_count == 1
    assert caption.caption.await_count == 1


def test_high_confidence_plain_ocr_does_not_call_vlm():
    ocr = MagicMock(spec=ImageRecognitionProvider)
    ocr.is_configured.return_value = True
    ocr.recognize_image = AsyncMock(
        return_value=RecognitionResult(text="普通文字", confidence=0.99, provider_name="ocr")
    )
    caption = MagicMock(spec=ImageCaptionProvider)
    caption.is_configured.return_value = True
    caption.caption = AsyncMock(return_value="不应调用")
    parser = ImageParser(recognition_provider=ocr, caption_provider=caption)

    parser.parse_asset(ImageAsset(content=b"plain-image", source_name="a.png"))

    caption.caption.assert_not_awaited()


def test_ocr_chart_layout_signal_triggers_vlm():
    ocr = MagicMock(spec=ImageRecognitionProvider)
    ocr.is_configured.return_value = True
    ocr.recognize_image = AsyncMock(
        return_value=RecognitionResult(
            text="quarterly data",
            confidence=0.99,
            provider_name="ocr",
            metadata={"chart_detected": True},
        )
    )
    caption = MagicMock(spec=ImageCaptionProvider)
    caption.is_configured.return_value = True
    caption.caption = AsyncMock(return_value="bar chart of quarterly data")
    parser = ImageParser(recognition_provider=ocr, caption_provider=caption)

    parser.parse_asset(ImageAsset(content=b"chart-image", source_name="report.png"))

    caption.caption.assert_awaited_once()


def test_identical_image_chunks_reuse_embedding_without_sharing_metadata():
    service = IndexingService.__new__(IndexingService)
    service.settings = SimpleNamespace(
        embedding_api_base="http://tei.internal/v1",
        embedding_model="BAAI/bge-m3",
        embedding_local_model=None,
        embedding_vector_size=1024,
    )
    service.embedding_provider = MagicMock()
    service.embedding_provider.embed_texts.return_value = [
        EmbeddingResult(
            text="same image description",
            vector=[0.1, 0.2],
            model_name="BAAI/bge-m3",
            provider_name="tei",
            request_tokens=3,
        )
    ]
    IndexingService._image_embedding_cache.clear()
    chunks = [
        SimpleNamespace(chunk_text="same image description", metadata_json={"image_hash": "image-sha"}),
        SimpleNamespace(chunk_text="same image description", metadata_json={"image_hash": "image-sha"}),
    ]

    first = service._embed_chunks(chunks)
    second = service._embed_chunks(chunks)

    assert len(first) == 2
    assert first[0].vector == first[1].vector
    assert first[0].vector is not first[1].vector
    assert second[0].vector == first[0].vector
    service.embedding_provider.embed_texts.assert_called_once_with(["same image description"])
