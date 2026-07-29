from __future__ import annotations

from pathlib import Path

from app.ingestion.image_assets import ImageAsset
from app.ingestion.types import ParsedBlock
from app.ingestion.parsers.common import normalize_text


class PptxParser:
    """Extract slide text and embedded raster images from PowerPoint files."""

    def parse(self, file_path: Path) -> list[ParsedBlock]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise RuntimeError("python-pptx is required to parse .pptx files") from exc

        presentation = Presentation(file_path)
        blocks: list[ParsedBlock] = []
        assets: list[ImageAsset] = []
        for slide_no, slide in enumerate(presentation.slides, start=1):
            slide_text = normalize_text(
                "\n".join(
                    shape.text
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text.strip()
                )
            )
            if slide_text:
                blocks.append(
                    ParsedBlock(
                        text=slide_text,
                        page_no=slide_no,
                        metadata={"parser": "python-pptx", "slide_no": slide_no},
                    )
                )

            for shape_index, shape in enumerate(slide.shapes, start=1):
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                image = shape.image
                content_type = image.content_type or "image/png"
                assets.append(
                    ImageAsset(
                        content=image.blob,
                        source_name=f"{file_path.name}#slide-{slide_no}-image-{shape_index}",
                        mime_type=content_type,
                        page_no=slide_no,
                        metadata={
                            "source_pptx": file_path.name,
                            "extraction_method": "python_pptx_picture",
                            "slide_no": slide_no,
                            "shape_index": shape_index,
                        },
                    )
                )

        if assets:
            from app.ingestion.parsers.image_parser import ImageParser

            try:
                blocks.extend(ImageParser().parse_assets(assets))
            except Exception:
                import logging

                logging.getLogger(__name__).exception("PPTX image extraction failed: %s", file_path)
        return blocks

